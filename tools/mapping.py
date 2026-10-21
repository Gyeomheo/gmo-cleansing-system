"""매핑 등록 — 처음 보는 원본 값을 표로 만들고, 표나 한 줄 명령으로 매핑 엑셀에 등록한다.

    python tools/mapping.py propose
        1_Input 원본에서 매핑 키에 없는 값 → Reports/매핑등록.xlsx
        (Spend 큰 순 · 지금 어떻게 클렌징되는지 · 비슷한 기존 규칙으로 표준값 제안)
        사람: 표준값 확인 · 고치고 '등록' 칸에 O
    python tools/mapping.py apply
        매핑등록.xlsx 에서 O 인 행을 매핑 엑셀에 추가(백업 후) → 등록 결과 미리보기
    python tools/mapping.py add MX "S26 FE" Smartphones "S26 Series" "S26 FE"
    python tools/mapping.py add CE "Washing Machine" Washer Bespoke NONE
    python tools/mapping.py add MEDIA "Pinterest" "Paid Social" Others Others
        원본 값 하나 → 표준 3칸. 키(C_Key/A_Key/F_Key)는 파이프라인과 같은 정규화로 자동 계산
    python tools/mapping.py test MX --cat Smartphones --series "S26 Series" --value "S26 FE"
        지금 매핑으로 그 값이 어떻게 클렌징되는지
    python tools/mapping.py retire MX "S24 Series" "Other S Series"
        세대가 밀리면 그 시리즈(표준2)로 가던 행을 전부 다른 시리즈로(백업 후). 제품명(표준3)은 그대로.
        마스터 · 취합본 과거 행까지 바꾸려면 다음 실행을 --remap 으로

원본 값은 MX = Products, CE = Product Category, MEDIA = Media Platform
(Media Type 1 이 Affiliate 면 Media Type 2).
"""
from __future__ import annotations

import argparse
import difflib
import logging
import re
import shutil
import tempfile
from datetime import datetime
from pathlib import Path

import pandas as pd

from common import REPORT_DIR, _read, config, is_locked
from src.converter import convert_excel_to_json, load_map_from_json
from src.pipeline import (fast_normalize_text, run_ce_product_cleansing,
                          run_cleansing_pipeline, sanitize_column_headers)

KINDS = {
    "MX": (config.PRODUCT_MX_EXCEL, config.PRODUCT_COLS_MAP_MX),
    "CE": (config.PRODUCT_CE_EXCEL, config.PRODUCT_COLS_MAP_CE),
    "MEDIA": (config.MEDIA_MAP_EXCEL, config.MEDIA_COLS_MAP),
}
SHEET = REPORT_DIR / "매핑등록.xlsx"
BACKUP = config.CONFIG_DIR / "_backup"
COLS = ["등록", "구분", "원본 값", "등록 키", "표준1", "표준2", "표준3", "제안 근거",
        "현재 결과", "행", "Spend(USD)", "법인", "원본 1", "원본 2", "원본 3"]
MARKS = {"O", "Y", "1", "○", "V"}


def norm_key(value: str) -> str:
    return str(fast_normalize_text(pd.Series([value])).iloc[0])


def load_rules(kind: str) -> pd.DataFrame:
    path, cols = KINDS[kind]
    df = pd.read_excel(path).fillna("")
    return df[[cols["key"]] + cols["std_cols"]].astype(str)


def clean_combos(kind: str, raw1, raw2, raw3) -> list[list[str]]:
    """원본 3칸 목록 → 지금 매핑 엑셀로 클렌징한 3칸 목록. run.py 와 같은 함수 · 같은 JSON 변환을 쓴다.

    MX/CE: (Product Category, Product Series, Products)
    MEDIA: (Media Type 1, Media Type 2, Media Platform)
    """
    kind = kind.upper()
    path, cols = KINDS[kind]
    raw = pd.DataFrame(dict(zip(cols["raw_cols"], [list(raw1), list(raw2), list(raw3)])))
    logging.disable(logging.CRITICAL)
    try:
        with tempfile.TemporaryDirectory() as td:
            js = Path(td) / "map.json"
            convert_excel_to_json(path, js, cols["key"], cols["std_cols"])
            df_map = load_map_from_json(js, cols)
        if kind == "MEDIA":
            out, _ = run_cleansing_pipeline(raw, df_map, cols, is_media=True)
        elif kind == "MX":
            out, _ = run_cleansing_pipeline(raw, df_map, cols)
        else:
            out, _ = run_ce_product_cleansing(raw, df_map, cols)
    finally:
        logging.disable(logging.NOTSET)
    return out[[f"{c}_cleaned" for c in cols["raw_cols"]]].astype(str).values.tolist()


# ---------------------------------------------------------------- propose
GENERIC = ("other", "multi", "none", "nan", "")


def _compact(s: pd.Series | str):
    if isinstance(s, str):
        return "".join(s.lower().split())
    return s.astype(str).str.lower().str.replace(r"\s+", "", regex=True)


def is_generic(value: str) -> bool:
    v = str(value).strip().lower()
    return v.startswith(GENERIC[:2]) or v in GENERIC


def suggest(kind: str, key: str, rules: pd.DataFrame, current: list[str]) -> tuple[list[str], str]:
    """표준값 제안. 위에서부터 처음 맞는 것:
    ① 띄어쓰기 · 대소문자만 다른 기존 키 (Tab S10 Fe ↔ Tab S10 FE, Watch 9 ↔ Watch9)
    ② 표준 이름에 그 값이 들어 있는 규칙 (Z Fold 8 → 'Z Fold8 & Z Flip8')
    ③ 지금 결과가 일반값(Other · Multi)이 아니면 지금 결과 그대로
    ④ 아주 비슷한 기존 키 (유사도 0.85 이상)
    """
    kcol = KINDS[kind][1]["key"]
    std = KINDS[kind][1]["std_cols"]
    ck = _compact(key)
    hit = rules[_compact(rules[kcol]) == ck]
    if not hit.empty:
        return [hit.iloc[0][c] for c in std], f"띄어쓰기·대소문자만 다른 키 '{hit.iloc[0][kcol]}'"
    if len(ck) >= 4:
        # 앞뒤가 영숫자가 아닐 때만 — 'zfold8' 은 'Z Fold8 & Z Flip8' 에는 맞고 'Z Fold8 Ultra' 에는 안 맞는다
        pat = rf"(?:^|[^a-z0-9]){re.escape(ck)}(?:$|[^a-z0-9])"
        for col in (std[2], std[1]):
            hit = rules[_compact(rules[col]).str.contains(pat, regex=True) & ~rules[col].map(is_generic)]
            if not hit.empty:
                row = hit.iloc[0]
                return [row[c] for c in std], f"표준 이름에 포함 '{row[col]}'"
    # 미디어는 못 찾으면 원본을 그대로 두므로 '지금 결과'가 표준값이라는 보장이 없다
    if kind != "MEDIA" and current and not is_generic(current[1]):
        return list(current), "지금 결과 그대로"
    close = difflib.get_close_matches(key, rules[kcol].unique().tolist(), n=1, cutoff=0.85)
    if close:
        row = rules[rules[kcol] == close[0]].iloc[0]
        return [row[c] for c in std], f"비슷한 키 '{close[0]}'"
    return ["", "", ""], ""


def _spend(df: pd.DataFrame) -> pd.Series:
    col = (next((c for c in df.columns if c.startswith("Media Spend (USD)")), None)
           or next((c for c in df.columns if c.startswith("Media Spend")), None))
    if not col:
        return pd.Series(0.0, index=df.index)
    return pd.to_numeric(df[col].astype(str).str.replace(",", ""), errors="coerce").fillna(0)


def scan_new_values() -> pd.DataFrame:
    """1_Input 원본에서 매핑 키에 없는 값(처음 보는 값)을 찾는다.

    파이프라인은 처음 보는 값도 분류 단위로 '구제'하면 Unmapped 리포트에 남기지 않는다
    (예: Z Fold9 → Smartphones / Multi). 그래서 원본에서 키가 있는지를 직접 본다.
    MX · MEDIA 는 키를 대소문자까지 그대로, CE 는 소문자로 맞춘다(각 파이프라인의 조인 기준).
    """
    rules = {k: load_rules(k) for k in KINDS}
    mx_keys = set(rules["MX"]["C_Key"])
    ce_keys = set(rules["CE"]["A_Key"].str.lower().str.strip())
    md_keys = set(rules["MEDIA"]["F_Key"])
    r = rules["MX"]
    mx_valid = set(r["A_Standard"] + "|" + r["B_Standard"] + "|" + r["C_Standard"])
    found = []
    for path in sorted(config.INPUT_DIR.glob("*.csv")):
        df = _read(path)
        if df is None:
            continue
        df = sanitize_column_headers(df).rename(columns={"Product Category2": "Product Category",
                                                         "Products (optional)": "Products"})
        need = ["Subsidiary", "Media Type 1", "Media Type 2", "Media Platform",
                "Product Category", "Product Series"]
        if any(c not in df.columns for c in need):
            continue
        df = df.fillna("")
        if "Products" not in df.columns:
            df["Products"] = ""
        kind = "CE" if config.COL_BU in df.columns else "MX"
        sp = _spend(df)
        sub = df["Subsidiary"].str.strip().str.upper()

        if kind == "MX":
            key = fast_normalize_text(df["Products"])
            combo = df["Product Category"] + "|" + df["Product Series"] + "|" + df["Products"]
            new = ~key.isin(mx_keys) & ~combo.isin(mx_valid)
            value = df["Products"]
        else:
            key = fast_normalize_text(df["Product Category"])
            new = ~key.str.lower().str.strip().isin(ce_keys)
            value = df["Product Category"]
        new &= ~key.isin(["", "-"])
        found.append(pd.DataFrame({
            "구분": kind, "원본 1": df["Product Category"], "원본 2": df["Product Series"],
            "원본 3": df["Products"], "원본 값": value, "등록 키": key, "법인": sub, "spend": sp})[new])

        aff = fast_normalize_text(df["Media Type 1"]) == "Affiliate"
        mvalue = df["Media Platform"].where(~aff, df["Media Type 2"])
        mkey = fast_normalize_text(mvalue)
        mnew = ~mkey.isin(md_keys) & (mkey != "")
        found.append(pd.DataFrame({
            "구분": "MEDIA", "원본 1": df["Media Type 1"], "원본 2": df["Media Type 2"],
            "원본 3": df["Media Platform"], "원본 값": mvalue, "등록 키": mkey, "법인": sub, "spend": sp})[mnew])

    allv = pd.concat(found, ignore_index=True) if found else pd.DataFrame()
    if allv.empty:
        return allv
    g = (allv.groupby(["구분", "원본 1", "원본 2", "원본 3", "원본 값", "등록 키"], sort=False)
         .agg(행=("spend", "size"), spend=("spend", "sum"),
              법인=("법인", lambda s: ", ".join(sorted(set(s)))))
         .reset_index())
    return g.sort_values("spend", ascending=False, ignore_index=True)


def propose() -> Path | None:
    g = scan_new_values()
    if not g.empty:
        g = g[g["구분"] != "MEDIA"]      # 미디어는 매핑을 늘리지 않는다 — 표준표 검사가 Others 로 처리
    if g.empty:
        print("■ 매핑 제안 — 처음 보는 값 없음")
        if SHEET.exists() and not is_locked(SHEET):
            SHEET.unlink()
        return None
    if is_locked(SHEET):
        print(f"  ✖ {SHEET.name} 이(가) 열려 있어 새로 못 씀 — 닫고 다시")
        return None
    rules = {k: load_rules(k) for k in KINDS}
    g["현재"] = None
    for kind in KINDS:
        m = g["구분"] == kind
        if m.any():
            g.loc[m, "현재"] = pd.Series(
                clean_combos(kind, g.loc[m, "원본 1"], g.loc[m, "원본 2"], g.loc[m, "원본 3"]),
                index=g.index[m])
    # 등록은 키 단위 → 키별로 한 줄. 대표 조합 = Spend 가 가장 큰 원본 조합
    rows = []
    for (kind, key), part in g.groupby(["구분", "등록 키"], sort=False):
        top = part.iloc[0]
        cur = list(top["현재"])
        std, why = suggest(kind, key, rules[kind], cur)
        if kind == "CE" and not std[2]:
            std[2] = "NONE"
        subs = ", ".join(sorted({s for x in part["법인"] for s in x.split(", ")}))
        note = f" (원본 조합 {len(part)}개)" if len(part) > 1 else ""
        rows.append(["", kind, top["원본 값"], key, *std, why, " / ".join(cur) + note,
                     int(part["행"].sum()), round(part["spend"].sum()), subs,
                     top["원본 1"], top["원본 2"], top["원본 3"], is_generic(cur[1])])
    out = pd.DataFrame(rows, columns=COLS + ["_generic"])
    # 지금 Other · Multi 로 빠지는 것부터(손볼 것), 그 안에서 Spend 큰 순
    out = out.sort_values(["_generic", "Spend(USD)"], ascending=[False, False]).drop(columns="_generic")
    with pd.ExcelWriter(SHEET, engine="openpyxl") as w:
        out.to_excel(w, index=False, sheet_name="등록")
        ws = w.sheets["등록"]
        for i, width in enumerate((6, 7, 24, 24, 14, 20, 20, 22, 44, 7, 11, 12, 16, 18, 24)):
            ws.column_dimensions[chr(65 + i)].width = width
        ws.freeze_panes = "C2"
    print(f"■ 매핑 제안 — 처음 보는 값 {len(out)}건(Spend 큰 순) → {SHEET}")
    print("  '현재 결과'가 맞으면 그냥 두고, 고칠 것만 표준1~3 을 채워 '등록' 칸에 O → 매핑등록.cmd")
    return SHEET


# ---------------------------------------------------------------- add / apply
def add_rules(items: list[tuple[str, str, list[str]]], notes: list[str] | None = None) -> int:
    """items = [(kind, 원본 값, [표준1, 표준2, 표준3])]. 엑셀별로 백업 후 한 번에 추가.
    notes 를 주면 매핑 엑셀 '메모' 열(없으면 만든다)에 같이 적는다 — 자동 등록 표시용."""
    from openpyxl import load_workbook

    added = 0
    by_kind: dict[str, list] = {}
    notes = notes or [""] * len(items)
    for (kind, value, std), note in zip(items, notes):
        by_kind.setdefault(kind.upper(), []).append((value, std, note))
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    for kind, entries in by_kind.items():
        path, cols = KINDS[kind]
        if is_locked(path):
            print(f"  ✖ {path.name} 이(가) Excel 에서 열려 있음 — 닫고 다시")
            continue
        rules = load_rules(kind)
        existing = set(map(tuple, rules.values.tolist()))
        new_rows = []
        for value, std, note in entries:
            key = norm_key(value)
            row = (key, *[str(s).strip() for s in std])
            if row in existing:
                print(f"  · {kind} '{value}' → 이미 같은 규칙 있음, 건너뜀")
                continue
            same_key = rules[rules[cols["key"]] == key]
            if not same_key.empty:
                print(f"  · {kind} 키 '{key}' 에 다른 규칙 {len(same_key)}개 있음"
                      " — 후보로 함께 둠(점수 높은 쪽이 쓰임)")
            new_rows.append((row, note))
            existing.add(row)
        if not new_rows:
            continue
        BACKUP.mkdir(exist_ok=True)
        shutil.copy2(path, BACKUP / f"{path.stem}_{stamp}{path.suffix}")
        wb = load_workbook(path)
        ws = wb.worksheets[0]
        header = [c.value for c in ws[1]]
        if any(n for _, n in new_rows) and "메모" not in header:
            ws.cell(row=1, column=len(header) + 1, value="메모")
            header.append("메모")
        order = [header.index(c) for c in [cols["key"]] + cols["std_cols"]]
        for row, note in new_rows:
            line = [None] * len(header)
            for pos, v in zip(order, row):
                line[pos] = v
            if note:
                line[header.index("메모")] = note
            ws.append(line)
            print(f"  ✔ {kind} 추가: {row[0]} → {' / '.join(row[1:])}")
        wb.save(path)
        added += len(new_rows)
    if added:
        print(f"  (백업: 3_Config/_backup/*_{stamp}.xlsx)")
    return added


def retire_series(kind: str, old: str, new: str) -> list[tuple[str, str]]:
    """표준2(시리즈)가 old 인 매핑 행을 모두 new 로 — 백업 후 서식 그대로."""
    from openpyxl import load_workbook

    path, cols = KINDS[kind]
    if is_locked(path):
        print(f"  ✖ {path.name} 이(가) Excel 에서 열려 있음 — 닫고 다시")
        return 0
    BACKUP.mkdir(exist_ok=True)
    shutil.copy2(path, BACKUP / f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}")
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    header = [c.value for c in ws[1]]
    ki, ci, si = (header.index(cols["key"]), header.index(cols["std_cols"][0]),
                  header.index(cols["std_cols"][1]))
    hits = []                                            # (키, 표준1)
    for row in ws.iter_rows(min_row=2):
        if str(row[si].value or "").strip() == old:
            row[si].value = new
            hits.append((str(row[ki].value), str(row[ci].value or "")))
    if not hits:
        print(f"  · {kind} 표준2 '{old}' 인 행 없음")
        return []
    wb.save(path)
    keys = [k for k, _ in hits]
    print(f"  ✔ {kind} {old} → {new} {len(keys)}행: {', '.join(keys[:12])}{' …' if len(keys) > 12 else ''}")
    return hits


def apply_sheet() -> int:
    if not SHEET.exists():
        print(f"  {SHEET.name} 없음 — 먼저 실행(또는 mapping.py propose)")
        return 0
    if is_locked(SHEET):
        print(f"  ✖ {SHEET.name} 이(가) 열려 있음 — 저장하고 닫은 뒤 다시")
        return 0
    df = pd.read_excel(SHEET, dtype=str).fillna("")
    pick = df[df["등록"].str.strip().str.upper().isin(MARKS)]
    if pick.empty:
        print("  '등록' 칸에 O 표시한 행이 없습니다.")
        return 0
    bad = pick[(pick["표준1"] == "") | (pick["표준2"] == "") | (pick["표준3"] == "")]
    if not bad.empty:
        print(f"  ✖ 표준값이 빈 행 {len(bad)}개 — 채우고 다시: " + ", ".join(bad["원본 값"]))
        return 0
    print("■ 매핑 등록")
    n = add_rules([(r["구분"], r["원본 값"], [r["표준1"], r["표준2"], r["표준3"]])
                   for _, r in pick.iterrows()])
    for kind, part in pick.groupby("구분"):
        res = clean_combos(kind, part["원본 1"], part["원본 2"], part["원본 3"])
        for (_, r), x in zip(part.iterrows(), res):
            ok = "✔" if x == [r["표준1"], r["표준2"], r["표준3"]] or kind == "CE" else "≠"
            print(f"  {ok} {kind} 미리보기: {r['원본 1']} / {r['원본 2']} / {r['원본 3']}  ⇒  {' / '.join(x)}")
    return n


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest="cmd", required=True)
    sp.add_parser("propose")
    sp.add_parser("apply")
    kinds = ["MX", "CE", "MEDIA", "mx", "ce", "media"]
    a_add = sp.add_parser("add")
    a_add.add_argument("kind", choices=kinds)
    a_add.add_argument("value")
    a_add.add_argument("std", nargs=3)
    a_test = sp.add_parser("test")
    a_test.add_argument("kind", choices=kinds)
    a_test.add_argument("--cat", "--type1", dest="a", default="")
    a_test.add_argument("--series", "--type2", dest="b", default="")
    a_test.add_argument("--value", "--platform", dest="c", default="")
    a_ret = sp.add_parser("retire")
    a_ret.add_argument("kind", choices=kinds)
    a_ret.add_argument("old")
    a_ret.add_argument("new")
    a = ap.parse_args(argv)
    if a.cmd == "retire":
        kind = a.kind.upper()
        hits = retire_series(kind, a.old, a.new)
        for key, cat in hits[:3]:
            raw = (cat, a.old, key) if kind != "CE" else (key, a.old, "")
            print(f"  ▷ 미리보기: {' / '.join(raw)}  ⇒  {' / '.join(clean_combos(kind, *[[v] for v in raw])[0])}")
        return 0 if hits else 1
    if a.cmd == "propose":
        propose()
    elif a.cmd == "apply":
        return 0 if apply_sheet() else 1
    elif a.cmd == "add":
        kind = a.kind.upper()
        if add_rules([(kind, a.value, a.std)]):
            raw = {"MX": (a.std[0], a.std[1], a.value), "CE": (a.value, a.std[1], ""),
                   "MEDIA": (a.std[0], a.std[1], a.value)}[kind]
            x = clean_combos(kind, [raw[0]], [raw[1]], [raw[2]])[0]
            print(f"  ▷ 미리보기: {' / '.join(raw)}  ⇒  {' / '.join(x)}")
    else:
        x = clean_combos(a.kind, [a.a], [a.b], [a.c])[0]
        print(f"  ▷ {a.kind.upper()}: {a.a} / {a.b} / {a.c}  ⇒  {' / '.join(x)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
