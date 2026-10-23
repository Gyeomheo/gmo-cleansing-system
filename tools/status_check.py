"""GMO 백과사전 'D2C MX Status Check' · 'D2C CE Status Check' 시트 기재.

    python tools/status_check.py 43 CE             미리보기(무엇을 적을지 표만)
    python tools/status_check.py 43 CE --apply     실제로 적기(백업 후 Excel 로 저장)
    python tools/status_check.py 43 CE --apply --book <사본.xlsx>   (시험용 — 받은 날짜는 공유 드라이브 폴더에서 읽기만)

- 받은 날짜: 국가별 weekly/<nn>주차/<MX|CE>/<법인 MMDD…> 폴더 이름의 날짜 → 그 날짜의 ISO 주차 칸에 MM-DD
  (예: CE 'SEAU 1001' → W40 칸에 10-01). 2026 칸 묶음(맨 뒤 W01~ · '정리' 앞)만. **빈 칸만 채운다** — 손으로 적은 값은 그대로
  · 그 주차 칸이 아직 없으면 '정리' 앞에 새 칸(Wnn)을 넣는다(왼쪽 칸 서식)
- 정리 · 취합: 매주 최신화 칸 — 다 지우고 이번 주에 클렌징 · 취합한 법인만 O
  (이번 주 법인 = 1_Input 의 그 구분 입력, 취합 O 는 그 주차 취합본이 있을 때)
- 서식: 매번 2026 주차 칸 전체를 MM-DD(`mm-dd`)로 — 손으로 적은 일련번호(46302 로 보이던 칸)도
- 법인 행: Subsidiary 칸이 같은 행 중 Media 칸이 비었거나 'X' 인 첫 행
- GUI 는 ③ 취합본 만들기 뒤에 자동 실행(tools/gui.py step_status)
"""
from __future__ import annotations

import argparse
import re
import shutil
from datetime import date, datetime
from pathlib import Path

import openpyxl

from common import config, current_slots
from collect import WEEKLY_ROOT

BOOK = WEEKLY_ROOT.parent / "4. Standard" / "★★GMO 백과사전_250327.xlsx"
SHEET = {"MX": "D2C MX Status Check", "CE": "D2C CE Status Check"}
HEAD_ROW, FIRST_ROW, SUB_COL, MEDIA_COL = 5, 6, 4, 5
YEAR = 2026


def received(kind: str) -> dict[str, dict[int, date]]:
    """법인 → {ISO 주차: 받은 날} (같은 주에 여러 번이면 마지막 날)."""
    out: dict[str, dict[int, date]] = {}
    for wdir in WEEKLY_ROOT.glob("*주차"):
        m = re.match(r"(\d+)주차$", wdir.name)
        if not m or not (wdir / kind).is_dir():
            continue
        wk = int(m.group(1))
        for f in (wdir / kind).iterdir():
            fm = re.match(r"\s*([A-Za-z]+(?:_[A-Za-z]+)*)[\s_]*(\d{4})", f.name)
            if not f.is_dir() or not fm:
                continue
            mm, dd = int(fm.group(2)[:2]), int(fm.group(2)[2:])
            try:
                d = date(YEAR - 1 if mm == 12 and wk <= 5 else YEAR, mm, dd)
            except ValueError:
                continue
            y, w, _ = d.isocalendar()
            if y != YEAR:
                continue
            cur = out.setdefault(fm.group(1).upper(), {})
            cur[w] = max(cur.get(w, d), d)
    return out


def _as_date(v) -> date | None:
    """시트 값 → 날짜(datetime · Excel 일련번호 둘 다 있음)."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        return date.fromordinal(date(1899, 12, 30).toordinal() + int(v))
    return None


def layout(ws) -> tuple[int, dict[int, int], dict[str, int]]:
    """('정리' 칸, {주차: 칸}, {법인: 행}) — 2026 묶음은 '정리' 바로 왼쪽의 W01~ 연속 칸."""
    head = {c: str(ws.cell(HEAD_ROW, c).value or "").strip() for c in range(1, ws.max_column + 1)}
    jeongri = next(c for c, v in head.items() if v == "정리")
    weeks, c = {}, jeongri - 1
    while c > 0 and re.fullmatch(r"W\d{2}", head[c]):
        weeks[int(head[c][1:])] = c
        if head[c] == "W01":
            break
        c -= 1
    rows: dict[str, int] = {}
    for r in range(FIRST_ROW, ws.max_row + 1):
        sub = str(ws.cell(r, SUB_COL).value or "").strip().upper()
        media = str(ws.cell(r, MEDIA_COL).value or "").strip().upper()
        if sub and sub not in rows and media in ("", "X"):
            rows[sub] = r
    return jeongri, weeks, rows


def plan(week: int, kind: str):
    wb = openpyxl.load_workbook(BOOK, read_only=False)
    ws = wb[SHEET[kind]]
    jeongri, weeks, rows = layout(ws)
    got = received(kind)
    need_weeks = sorted({w for d in got.values() for w in d} - set(weeks))
    add_cols = [w for w in range(max(weeks) + 1, max(need_weeks, default=0) + 1)]
    col_of = dict(weeks)
    for i, w in enumerate(add_cols):
        col_of[w] = jeongri + i                          # 새 칸은 '정리' 자리부터 → '정리'는 오른쪽으로
    shift = len(add_cols)
    fills, skipped = [], []
    no_row = sorted(s for s in got if s not in rows and not any(k.startswith(s + "_") for k in rows))
    for sub, by_week in sorted(got.items()):
        # 폴더 'SEBN' · 'SECZ' = 시트 SEBN_BE · SEBN_NL / SECZ_CZ · SECZ_SK 모두
        targets = [rows[sub]] if sub in rows else [r for s, r in rows.items() if s.startswith(sub + "_")]
        for r in targets:
            for w, d in sorted(by_week.items()):
                old = ws.cell(r, weeks[w]).value if w in weeks else None
                if old in (None, ""):
                    fills.append((ws.cell(r, SUB_COL).value, r, w, col_of[w], d))
                elif _as_date(old) != d:
                    skipped.append((ws.cell(r, SUB_COL).value, w, old, d))
    return dict(ws=ws, w01=weeks[min(weeks)], jeongri=jeongri, shift=shift, add_cols=add_cols, fills=fills, skipped=skipped,
                no_row=no_row, rows=rows, last_row=ws.max_row)


def _this_week_subs(kind: str) -> list[str]:
    subs = set()
    for infos in current_slots().values():
        for i in infos:
            if i.kind == kind and i.sub:
                subs.add(i.sub.upper())
    return sorted(subs)


def _consolidated(week: int, kind: str) -> bool:
    from consolidate import NAME_RE, latest_file
    f = latest_file(kind)
    m = NAME_RE.match(f.name) if f else None
    return bool(m and int(m.group(2)) == week)


def run(week: int, kind: str, apply: bool) -> int:
    p = plan(week, kind)
    this = _this_week_subs(kind)
    consol = _consolidated(week, kind)
    print(f"■ STATUS 기재 — {SHEET[kind]} ({BOOK.name})")
    if p["add_cols"]:
        print(f"  새 주차 칸: {', '.join(f'W{w:02d}' for w in p['add_cols'])} ('정리' 앞)")
    print(f"  받은 날짜 채움 {len(p['fills'])}칸")
    for sub, _, w, _, d in p["fills"]:
        print(f"    {sub:<8} W{w:02d}  {d:%m-%d}")
    if p["skipped"]:
        print(f"  이미 다른 값이 있어 그대로 {len(p['skipped'])}칸")
        for sub, w, old, d in p["skipped"]:
            o = _as_date(old).strftime("%m-%d") if _as_date(old) else old
            print(f"    {sub:<8} W{w:02d}  시트 {o} · 폴더 {d:%m-%d}")
    if p["no_row"]:
        print(f"  ⚠ 시트에 행이 없는 법인(건너뜀): {', '.join(p['no_row'])}")
    missing = [s for s in this if s not in p["rows"]]
    if not this:
        print(f"  정리 · 취합: 이번 주 {kind} 입력 없음 — 그대로")
    else:
        print(f"  정리 · 취합: 지우고 → 정리 O {', '.join(s for s in this if s in p['rows']) or '없음'}"
              f" · 취합 {'O (같은 법인)' if consol else '— (W' + str(week) + ' 취합본 없음)'}")
    if missing:
        print(f"  ⚠ 정리 O 를 적을 행이 없는 법인: {', '.join(missing)}")
    if not apply:
        print("\n  미리보기만 — 적으려면 --apply")
        return 0

    backup = config.BACKUP_DIR / f"백과사전_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    config.BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(BOOK, backup)
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.DisplayAlerts = False
    xl.ScreenUpdating = False
    try:
        wb = xl.Workbooks.Open(str(BOOK))
        try:
            if wb.ReadOnly:
                print("  ✖ 백과사전이 읽기 전용으로 열림(누가 열어 둠?) — 닫고 다시")
                return 1
            ws = wb.Worksheets(SHEET[kind])
            j = p["jeongri"]
            for i, w in enumerate(p["add_cols"]):
                ws.Columns(j + i).Insert(Shift=-4161, CopyOrigin=0)     # xlToRight, 왼쪽 칸 서식
                ws.Range(ws.Cells(HEAD_ROW + 1, j + i), ws.Cells(p["last_row"], j + i)).ClearContents()
                ws.Cells(HEAD_ROW, j + i).Value = f"W{w:02d}"
            for sub, r, w, c, d in p["fills"]:
                cell = ws.Cells(r, c)
                cell.Value2 = (d - date(1899, 12, 30)).days            # Excel 일련번호(COM datetime 은 시간대로 밀림)
            # 서식: 2026 주차 칸 전체(손 입력 포함) MM-DD — 일련번호가 숫자(46302)로 보이던 칸까지
            # 열마다 지정 — 숨긴 열이 섞인 범위에 한 번에 넣으면 Excel 이 일부 칸을 말없이 건너뛴다
            for c in range(p["w01"], j + p["shift"]):
                ws.Range(ws.Cells(FIRST_ROW, c), ws.Cells(p["last_row"], c)).NumberFormat = "mm-dd"
            block = ws.Range(ws.Cells(FIRST_ROW, p["w01"]), ws.Cells(p["last_row"], j + p["shift"] - 1))
            jc, cc = j + p["shift"], j + p["shift"] + 1                 # 정리 · 취합
            if this:                                                    # 이번 주 그 구분 입력이 없으면 그대로
                ws.Range(ws.Cells(FIRST_ROW, jc), ws.Cells(p["last_row"], cc)).ClearContents()
            for s in this:
                r = p["rows"].get(s)
                if r:
                    ws.Cells(r, jc).Value = "O"
                    if consol:
                        ws.Cells(r, cc).Value = "O"
            wb.Save()
            nums = [c for c in block if isinstance(c.Value2, float)]
            odd = [c.Address for c in nums if c.NumberFormat != "mm-dd"]
            print(f"  서식 확인: 2026 날짜 {len(nums)}칸 중 MM-DD 아닌 칸 {len(odd)}개" + (f" — {odd[:5]}" if odd else ""))
        finally:
            wb.Close(False)
    finally:
        xl.Quit()
    print(f"\n  ✔ 저장 · 백업 {backup.relative_to(config.OUTPUT_DIR.parent)}")
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("week", type=int)
    ap.add_argument("kind", choices=["CE", "MX", "ce", "mx"])
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--book", type=Path, help="백과사전 대신 이 파일(시험용 사본)")
    a = ap.parse_args(argv)
    if a.book:
        global BOOK
        BOOK = a.book
    return run(a.week, a.kind.upper(), a.apply)


if __name__ == "__main__":
    raise SystemExit(main())
